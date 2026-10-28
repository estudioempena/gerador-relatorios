import io
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageOps

PASTA_SISTEMA = Path(__file__).resolve().parent
PASTA_PROJETO = PASTA_SISTEMA.parent
PASTA_TEMPLATES = PASTA_PROJETO / "Templates"
PASTA_SAIDA = PASTA_PROJETO / "Relatorios_Gerados"
PASTA_FONTES = PASTA_SISTEMA / "fonts"

FONTE_TITULO = "Clash Display Semibold"
FONTE_CORPO = "Clash Display Regular"
COR_PRETO = RGBColor(0, 0, 0)

LARGURA_COLUNA_CM = 8.2
ALTURA_MAX_FOTO_CM = 6.0


def registrar_fontes_linux() -> None:
    # Garante que servidores Linux na nuvem reconheçam a fonte Clash Display
    if sys.platform != "win32" and PASTA_FONTES.exists():
        pasta_fontes_user = Path.home() / ".local" / "share" / "fonts"
        pasta_fontes_user.mkdir(parents=True, exist_ok=True)
        for f in PASTA_FONTES.glob("*.otf"):
            destino = pasta_fontes_user / f.name
            if not destino.exists():
                shutil.copy(f, destino)
        subprocess.run(["fc-cache", "-f"], capture_output=True)


def encontrar_template() -> Path:
    templates = [t for t in PASTA_TEMPLATES.glob("*.docx") if not t.name.startswith("~$")]
    if not templates:
        raise FileNotFoundError(f"Nenhum template .docx encontrado em: {PASTA_TEMPLATES}")
    return templates[0]


def aplicar_fonte(run, nome_fonte: str, tamanho_pt: float, cor: RGBColor = COR_PRETO):
    run.font.name = nome_fonte
    run.font.size = Pt(tamanho_pt)
    run.font.color.rgb = cor
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), nome_fonte)
    rFonts.set(qn("w:hAnsi"), nome_fonte)
    rFonts.set(qn("w:cs"), nome_fonte)


def ajustar_margens_e_cabecalho(doc: Document) -> None:
    secao = doc.sections[0]
    secao.left_margin = Cm(2.0)
    secao.right_margin = Cm(2.0)

    header = secao.header
    for drawing in header._element.xpath(".//*[local-name()='drawing']"):
        docPr = drawing.xpath(".//*[local-name()='docPr']")
        if docPr and "Caixa de Texto 5" in docPr[0].get("name", ""):
            posH = drawing.xpath(".//*[local-name()='positionH']")
            if posH:
                posH[0].set("relativeFrom", "page")
                offset = posH[0].xpath(".//*[local-name()='posOffset']")
                if offset:
                    offset[0].text = "644390"


def atualizar_data_cabecalho(doc: Document, data_texto: str) -> None:
    header = doc.sections[0].header
    for p in header._element.xpath(".//w:p[not(.//w:p)]"):
        t_nodes = p.xpath(".//w:t")
        texto = "".join(t.text for t in t_nodes if t.text)
        if "20XX" in texto:
            runs = p.xpath(".//w:r")
            if runs:
                primeiro_t = runs[0].xpath(".//w:t")
                if primeiro_t:
                    primeiro_t[0].text = data_texto
                for r in runs[1:]:
                    p.remove(r)


def normalizar_orientacao_imagem(caminho_foto: Path) -> Path:
    with Image.open(caminho_foto) as img:
        img_corrigida = ImageOps.exif_transpose(img)
        if img.size != img_corrigida.size or getattr(img, "_getexif", lambda: None)():
            buffer = io.BytesIO()
            formato = img.format or "JPEG"
            if formato == "JPEG" and img_corrigida.mode in ("RGBA", "P"):
                img_corrigida = img_corrigida.convert("RGB")
            img_corrigida.save(buffer, format=formato, quality=95)
            with open(caminho_foto, "wb") as f_out:
                f_out.write(buffer.getvalue())
    return caminho_foto


def calcular_dimensoes_foto(caminho_foto: Path) -> tuple[Cm, Cm]:
    normalizar_orientacao_imagem(caminho_foto)
    with Image.open(caminho_foto) as img:
        largura_px, altura_px = img.size
    escala = min(LARGURA_COLUNA_CM / largura_px, ALTURA_MAX_FOTO_CM / altura_px)
    return Cm(largura_px * escala), Cm(altura_px * escala)


def inserir_grade_fotos(doc: Document, fotos: List[Dict[str, str]]) -> None:
    for i in range(0, len(fotos), 2):
        par = fotos[i : i + 2]
        tabela = doc.add_table(rows=1, cols=2)
        tabela.alignment = WD_TABLE_ALIGNMENT.LEFT

        for col_idx, item in enumerate(par):
            caminho_img = Path(item["caminho"])
            legenda = item.get("legenda", "").strip()
            cell = tabela.rows[0].cells[col_idx]
            cell.width = Cm(LARGURA_COLUNA_CM)

            largura, altura = calcular_dimensoes_foto(caminho_img)
            p_img = cell.paragraphs[0]
            p_img.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_img.paragraph_format.left_indent = Cm(0)
            p_img.paragraph_format.space_before = Pt(0)
            p_img.paragraph_format.space_after = Pt(2)
            run_img = p_img.add_run()
            run_img.add_picture(str(caminho_img), width=largura, height=altura)

            if legenda:
                p_leg = cell.add_paragraph()
                p_leg.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_leg.paragraph_format.left_indent = Cm(0)
                p_leg.paragraph_format.space_before = Pt(0)
                p_leg.paragraph_format.space_after = Pt(6)
                p_leg.paragraph_format.line_spacing = 1.0
                run_leg = p_leg.add_run(legenda)
                aplicar_fonte(run_leg, FONTE_CORPO, 8.5)

        if len(par) == 1:
            tabela.rows[0].cells[1].width = Cm(LARGURA_COLUNA_CM)


def salvar_com_fallback(doc: Document, caminho_desejado: Path) -> Path:
    PASTA_SAIDA.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(caminho_desejado)
        return caminho_desejado
    except PermissionError:
        timestamp = datetime.now().strftime("%H%M%S")
        caminho_alternativo = caminho_desejado.with_name(
            f"{caminho_desejado.stem}_{timestamp}{caminho_desejado.suffix}"
        )
        doc.save(caminho_alternativo)
        return caminho_alternativo


def converter_docx_para_pdf(caminho_docx: Path, caminho_pdf: Path) -> Path:
    caminho_final_pdf = caminho_pdf
    if caminho_final_pdf.exists():
        try:
            with open(caminho_final_pdf, "a+"):
                pass
        except PermissionError:
            timestamp = datetime.now().strftime("%H%M%S")
            caminho_final_pdf = caminho_final_pdf.with_name(
                f"{caminho_pdf.stem}_{timestamp}{caminho_pdf.suffix}"
            )

    # 1. MODO WINDOWS (execução local via Microsoft Word)
    if sys.platform == "win32":
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        word = None
        doc = None
        try:
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0

            doc = word.Documents.Open(str(caminho_docx.resolve()))
            wdFormatPDF = 17
            doc.SaveAs(str(caminho_final_pdf.resolve()), FileFormat=wdFormatPDF)
            return caminho_final_pdf
        finally:
            if doc is not None:
                doc.Close(0)
            if word is not None:
                word.Quit()
            pythoncom.CoUninitialize()

    # 2. MODO NUVEM / LINUX (execução online no Streamlit Cloud via LibreOffice)
    else:
        registrar_fontes_linux()
        cmd = [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(caminho_final_pdf.parent.resolve()),
            str(caminho_docx.resolve()),
        ]
        resultado = subprocess.run(cmd, capture_output=True, text=True)
        if resultado.returncode != 0:
            raise RuntimeError(f"Erro na conversão LibreOffice: {resultado.stderr}")
        return caminho_final_pdf


def gerar_relatorio(
    titulo_geral: str,
    blocos: List[Dict[str, Any]],
    data_documento: Optional[str] = None,
    nome_base: str = "relatorio_final",
    gerar_pdf: bool = True,
) -> tuple[Path, Optional[Path]]:
    doc = Document(encontrar_template())

    ajustar_margens_e_cabecalho(doc)

    data_formatada = data_documento or datetime.now().strftime("%d . %m . %Y")
    atualizar_data_cabecalho(doc, data_formatada)

    # Título Geral do Documento
    p_titulo = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_titulo.paragraph_format.left_indent = Cm(0)
    p_titulo.paragraph_format.space_before = Pt(0)
    p_titulo.paragraph_format.space_after = Pt(10)
    run_titulo = p_titulo.add_run(titulo_geral)
    aplicar_fonte(run_titulo, FONTE_TITULO, 10)

    for bloco in blocos:
        subtitulo = bloco.get("titulo", "").strip()
        texto = bloco.get("texto", "").strip()
        fotos = bloco.get("fotos", [])

        if subtitulo:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_sub.paragraph_format.left_indent = Cm(0)
            p_sub.paragraph_format.space_before = Pt(12)
            p_sub.paragraph_format.space_after = Pt(4)
            run_sub = p_sub.add_run(subtitulo)
            aplicar_fonte(run_sub, FONTE_TITULO, 10)

        if texto:
            for linha in texto.split("\n\n"):
                linha_limpa = linha.strip()
                if not linha_limpa:
                    continue
                p_txt = doc.add_paragraph()
                p_txt.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_txt.paragraph_format.left_indent = Cm(0)
                p_txt.paragraph_format.space_before = Pt(0)
                p_txt.paragraph_format.space_after = Pt(8)
                p_txt.paragraph_format.line_spacing = 1.5
                run_txt = p_txt.add_run(linha_limpa)
                aplicar_fonte(run_txt, FONTE_CORPO, 10)

        if fotos:
            inserir_grade_fotos(doc, fotos)

    caminho_docx = salvar_com_fallback(doc, PASTA_SAIDA / f"{nome_base}.docx")

    caminho_pdf = None
    if gerar_pdf:
        caminho_pdf_alvo = caminho_docx.with_suffix(".pdf")
        try:
            caminho_pdf = converter_docx_para_pdf(caminho_docx, caminho_pdf_alvo)
        except Exception as e:
            print(f"Erro na conversão para PDF: {e}")
            raise RuntimeError(f"Falha ao gerar o PDF: {e}")

    return caminho_docx, caminho_pdf
